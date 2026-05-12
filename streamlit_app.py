from pathlib import Path
import base64

import openpyxl
import pandas as pd
import streamlit as st

from app import (
    arquivo_modelo_padrao,
    exit_dir,
    temp_dir,
    ABREVS,
    ABREV_PARA_MES,
    gerar_relatorio,
)

# ── Helpers ────────────────────────────────────────────────────────────────────
def carregar_imagem_base64(caminho: str) -> str:
    try:
        with open(caminho, "rb") as f:
            return base64.b64encode(f.read()).decode()
    except FileNotFoundError:
        return ""


def salvar_upload(uploaded_file, pasta: Path) -> Path:
    pasta.mkdir(exist_ok=True, parents=True)
    destino = pasta / uploaded_file.name
    with open(destino, "wb") as f:
        f.write(uploaded_file.getbuffer())
    return destino


def carregar_preview(caminho: Path) -> tuple[pd.DataFrame, str]:
    wb = openpyxl.load_workbook(caminho, data_only=True)
    nome_aba = next((n for n in wb.sheetnames if "49" in str(n)), wb.sheetnames[0])
    ws = wb[nome_aba]

    matriz = []
    for r in range(1, ws.max_row + 1):
        linha = [ws.cell(r, c).value for c in range(1, ws.max_column + 1)]
        matriz.append(linha)

    # expande células mescladas
    for mr in ws.merged_cells.ranges:
        val = ws.cell(mr.min_row, mr.min_col).value
        for r in range(mr.min_row, mr.max_row + 1):
            for c in range(mr.min_col, mr.max_col + 1):
                matriz[r - 1][c - 1] = val

    df = pd.DataFrame(matriz)
    df = df.loc[:, ~df.apply(lambda col: col.astype(str).str.strip().eq("").all())]
    df = df.loc[~df.apply(lambda row: row.astype(str).str.strip().eq("").all(), axis=1)]
    return df.reset_index(drop=True), nome_aba


def fmt_br(valor) -> str:
    """Formata número para padrão brasileiro."""
    if valor is None or (isinstance(valor, float) and pd.isna(valor)) or valor == "":
        return ""
    if isinstance(valor, (int, float)):
        return f"{valor:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")
    return str(valor)


def df_para_html(df: pd.DataFrame) -> str:
    if len(df) < 3:
        return "<p>Tabela vazia</p>"

    cabecalho = [str(v) if v is not None else "" for v in df.iloc[2].tolist()]
    corpo = df.iloc[3:]

    linhas = ["<thead><tr>"]
    for v in cabecalho:
        linhas.append(f"<th>{v}</th>")
    linhas.append("</tr></thead><tbody>")

    for _, row in corpo.iterrows():
        linhas.append("<tr>")
        for v in row:
            linhas.append(f"<td>{fmt_br(v)}</td>")
        linhas.append("</tr>")
    linhas.append("</tbody>")

    return (
        '<div class="preview-wrapper">'
        '<table class="preview-table">'
        + "".join(linhas)
        + "</table></div>"
    )


# ── CSS ────────────────────────────────────────────────────────────────────────
STYLE = """
<style>
.stApp { background-color: #d1d5db; }
section[data-testid="stSidebar"] {
    background-color: #fdf2f8 !important;
    border-right: 2px solid #fce7f3;
}
.stButton > button {
    background-color: #ff007f !important; color: white !important;
    border-radius: 12px !important; width: 100% !important;
    font-weight: 800 !important; padding: 15px !important; border: none !important;
}
.stDownloadButton > button {
    background-color: #831843 !important; color: white !important;
    border-radius: 12px !important; width: 100% !important;
    font-weight: 800 !important; padding: 15px !important; border: none !important;
}
div[data-testid="stFileUploaderDropzoneInstructions"],
div[data-testid="stFileUploaderDropzone"] small,
div[data-testid="stFileUploaderDropzone"] p { display: none !important; }
.preview-wrapper {
    background: white; border-radius: 12px; padding: 14px;
    box-shadow: 0 8px 18px rgba(0,0,0,0.08);
    overflow-x: auto; width: 100%;
}
.preview-table {
    border-collapse: collapse; font-size: 12px; background: white;
    table-layout: fixed; min-width: 1100px;
}
.preview-table th, .preview-table td {
    border: 1px solid #d1d5db; padding: 6px 8px;
    max-width: 140px; min-width: 90px;
    overflow: hidden; text-overflow: ellipsis; white-space: nowrap;
    color: #111827 !important;
}
.preview-table td { background-color: #ffffff !important; }
.preview-table tr:nth-child(even) td { background-color: #f9fafb !important; }
.preview-table th {
    background-color: #fdf2f8; color: #831843;
    font-weight: 700; position: sticky; top: 0; z-index: 1;
}
</style>
"""

# ── App ────────────────────────────────────────────────────────────────────────
img_b64 = carregar_imagem_base64(r"img\logo.png")

st.set_page_config(page_title="Monitor de Metas Institucionais", layout="wide")
st.markdown(STYLE, unsafe_allow_html=True)

# Cabeçalho
st.markdown(f"""
<div style="display:flex;align-items:center;justify-content:center;gap:15px;
            padding:20px;background:linear-gradient(90deg,#ff007f,#ff69b4);
            border-radius:12px;margin-bottom:25px;">
  {'<img src="data:image/png;base64,' + img_b64 + '" width="120">' if img_b64 else ''}
  <div>
    <h1 style="margin:0;color:white;">Monitor de Metas Municipais</h1>
    <p style="margin:0;opacity:0.7;color:white;">Coplan-Dados</p>
  </div>
</div>
""", unsafe_allow_html=True)

# ── Sidebar ────────────────────────────────────────────────────────────────────
with st.sidebar:
    st.markdown("<h2 style='color:#ff007f;text-align:center;'>Painel</h2>", unsafe_allow_html=True)

    ano_sel = st.selectbox("Selecione o ano", options=[2026, 2025], index=0)

    st.divider()
    st.markdown("**Arquivos de entrada**")

    word_file  = st.file_uploader("Word CONVIAS (.docx)", type=["docx"])
    excel_cons = st.file_uploader("Excel CONSEMAVI — Recapeamento (.xlsx)", type=["xlsx"])

    st.divider()
    btn_gerar = st.button("⚙️ Gerar relatório")

# ── Geração ────────────────────────────────────────────────────────────────────
if btn_gerar:
    erros = []
    if not word_file:
        erros.append("Word CONVIAS não enviado.")
    if not excel_cons:
        erros.append("Excel CONSEMAVI não enviado.")
    if not arquivo_modelo_padrao.exists():
        erros.append(f"Modelo não encontrado: coloque o arquivo `Acompanhamento_mensal_PDM_-_2026.xlsx` dentro da pasta `data/`.")

    if erros:
        for e in erros:
            st.error(e)
    else:
        with st.spinner("Lendo arquivos e preenchendo o modelo..."):
            try:
                caminho_word  = salvar_upload(word_file,  temp_dir)
                caminho_cons  = salvar_upload(excel_cons, temp_dir)

                caminho_saida, abrev_mes = gerar_relatorio(
                    caminho_word=caminho_word,
                    caminho_consemavi=caminho_cons,
                    ano_ref=ano_sel,
                )

                mes_nome = ABREV_PARA_MES[abrev_mes].capitalize()
                st.success(f"Relatório gerado! Mês detectado: **{mes_nome}/{ano_sel}**")

                st.session_state["arquivo_gerado"] = str(caminho_saida)
                st.session_state["mes_gerado"]     = mes_nome
                st.session_state["ano_gerado"]     = ano_sel

            except Exception as exc:
                st.error(f"Erro ao gerar relatório: {exc}")

# ── Botão de download + preview ───────────────────────────────────────────────
if "arquivo_gerado" in st.session_state:
    caminho_prev = Path(st.session_state["arquivo_gerado"])
    mes_g = st.session_state.get("mes_gerado", "")
    ano_g = st.session_state.get("ano_gerado", "")

    if caminho_prev.exists():
        with open(caminho_prev, "rb") as f:
            st.download_button(
                label="⬇️ Baixar Excel gerado",
                data=f,
                file_name=caminho_prev.name,
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )

        try:
            df_prev, aba = carregar_preview(caminho_prev)
            st.markdown(f"### Preview — {mes_g}/{ano_g} · aba `{aba}`")
            st.markdown(df_para_html(df_prev), unsafe_allow_html=True)
        except Exception as exc:
            st.warning(f"Não foi possível exibir o preview: {exc}")