from pathlib import Path
from typing import Optional
import re
import unicodedata

import openpyxl
import pandas as pd
from docx import Document
from openpyxl.cell.cell import MergedCell

# ── Paths ──────────────────────────────────────────────────────────────────────
base_dir = Path(__file__).resolve().parent
data_dir = base_dir / "data"
exit_dir = base_dir / "exit"
temp_dir = base_dir / "temp_inputs"
exit_dir.mkdir(exist_ok=True)
temp_dir.mkdir(exist_ok=True)

arquivo_modelo_padrao = data_dir / "Acompanhamento_mensal_PDM_-_2026.xlsx"

# ── Meses ──────────────────────────────────────────────────────────────────────
# Abreviação do modelo → nome completo
ABREV_PARA_MES = {
    "jan": "janeiro", "fev": "fevereiro", "mar": "março",
    "abr": "abril",   "mai": "maio",      "jun": "junho",
    "jul": "julho",   "ago": "agosto",    "set": "setembro",
    "out": "outubro", "nov": "novembro",  "dez": "dezembro",
}
MES_PARA_ABREV = {v: k for k, v in ABREV_PARA_MES.items()}
MESES = list(ABREV_PARA_MES.values())          # lista de nomes completos
ABREVS = list(ABREV_PARA_MES.keys())           # lista de abreviações

# Colunas do modelo (índice 0-based): col 3 = Jan … col 14 = Dez, col 15 = Total
ABREV_PARA_COL = {a: 3 + i for i, a in enumerate(ABREVS)}
COL_TOTAL = 15

# ── Normalização ───────────────────────────────────────────────────────────────
def norm(texto) -> str:
    """Remove acentos, caixa baixa, espaços extras."""
    if texto is None or (isinstance(texto, float) and pd.isna(texto)):
        return ""
    s = str(texto).strip().lower()
    s = unicodedata.normalize("NFKD", s).encode("ASCII", "ignore").decode("ASCII")
    return re.sub(r"\s+", " ", s)


def norm_sigla(s) -> str:
    s = norm(s)
    s = re.sub(r"\s+", "", s)
    # alias histórico fo → fb
    return "fb" if s == "fo" else s


def to_float(valor) -> float:
    if valor is None or (isinstance(valor, float) and pd.isna(valor)):
        return 0.0
    if isinstance(valor, (int, float)):
        return float(valor)
    t = str(valor).strip().replace("m²", "").replace("M²", "").replace("m2", "").strip()
    if t in {"", "-", "–", "—"}:
        return 0.0
    if "," in t:
        t = t.replace(".", "").replace(",", ".")
    t = re.sub(r"[^0-9.\-]", "", t)
    try:
        return float(t)
    except ValueError:
        return 0.0


def abrev_do_texto(texto: str) -> Optional[str]:
    """
    Extrai a abreviação do mês de strings como:
      'Mar/2026*', 'Março', 'março', 'mar', 'MAR'
    Retorna a abreviação minúscula (ex: 'mar') ou None.
    """
    t = norm(texto)
    # tenta abreviação direta (jan, fev, mar…)
    for a in ABREVS:
        if t == a or t.startswith(a + "/") or t.startswith(a + " "):
            return a
    # tenta nome completo
    for nome, a in MES_PARA_ABREV.items():
        if t == norm(nome):
            return a
    # tenta início do nome (marco → março)
    for nome, a in MES_PARA_ABREV.items():
        if norm(nome).startswith(t[:3]):
            return a
    return None


# ── Leitura do DOCX (CONVIAS) ─────────────────────────────────────────────────
def ler_word_convias(caminho_word: Path) -> tuple[pd.DataFrame, str]:
    """
    Lê o Word do CONVIAS.

    Formato esperado: tabela com cabeçalho contendo colunas
      Subprefeitura | Sigla | <Mês>/<Ano>[*]
    As linhas de dados têm: nome_subprefeitura | sigla | valor

    Retorna (DataFrame com colunas [sigla, convias], abrev_mes).
    """
    doc = Document(caminho_word)

    for tbl in doc.tables:
        if not tbl.rows:
            continue
        header = [cell.text.strip() for cell in tbl.rows[0].cells]

        # localiza coluna do mês (header com padrão abrev/ano ou nome do mês)
        col_mes_idx = None
        abrev_mes = None
        for idx, h in enumerate(header):
            a = abrev_do_texto(h)
            if a:
                col_mes_idx = idx
                abrev_mes = a
                break

        if col_mes_idx is None:
            continue  # esta tabela não tem mês, tenta a próxima

        # localiza coluna da sigla
        col_sigla_idx = None
        for idx, h in enumerate(header):
            if norm(h) == "sigla":
                col_sigla_idx = idx
                break

        if col_sigla_idx is None:
            continue

        dados = []
        for row in tbl.rows[1:]:
            cells = [cell.text.strip() for cell in row.cells]
            if len(cells) <= max(col_sigla_idx, col_mes_idx):
                continue
            sigla = norm_sigla(cells[col_sigla_idx])
            valor = to_float(cells[col_mes_idx])
            if sigla and re.fullmatch(r"[a-z]{1,3}", sigla):
                dados.append({"sigla": sigla, "convias": valor})

        if dados:
            df = pd.DataFrame(dados).groupby("sigla", as_index=False)["convias"].sum()
            return df, abrev_mes

    raise ValueError(
        "Não foi possível identificar a tabela do CONVIAS no Word.\n"
        "Verifique se o arquivo possui uma tabela com colunas 'Sigla' e '<Mês>/<Ano>'."
    )


# ── Leitura do Excel CONSEMAVI (Recapeamento) ─────────────────────────────────
def ler_excel_consemavi(caminho_excel: Path, ano_ref: int, abrev_mes: str) -> pd.DataFrame:
    """
    Lê o Excel do CONSEMAVI e retorna DataFrame [sigla, consemavi]
    para o ano e mês solicitados.
    """
    df_raw = pd.read_excel(caminho_excel, sheet_name="49", header=None)
    col2 = df_raw.iloc[:, 2].astype(str).apply(norm)

    titulo_bloco = f"area recapeada em {ano_ref}"
    inicios = col2[col2 == titulo_bloco].index.tolist()
    if not inicios:
        raise ValueError(f"Bloco '{titulo_bloco}' não encontrado no Excel do CONSEMAVI.")

    linha_titulo = inicios[0]

    # cabeçalho: linha com "sub" e pelo menos um mês
    linha_header = None
    for idx in range(linha_titulo, min(linha_titulo + 6, len(df_raw))):
        vals = [norm(x) for x in df_raw.iloc[idx].tolist()]
        if "sub" in vals and any(norm(m) in vals for m in MESES):
            linha_header = idx
            break
    if linha_header is None:
        raise ValueError("Cabeçalho do CONSEMAVI não encontrado.")

    # fim do bloco: primeiro "total" após o início
    fins = [i for i in col2[col2.str.contains("total", na=False)].index if i > linha_titulo]
    if not fins:
        raise ValueError("Linha de total do CONSEMAVI não encontrada.")
    linha_fim = fins[0]

    df_bloco = df_raw.iloc[linha_header + 1:linha_fim].copy()
    df_bloco.columns = df_raw.iloc[linha_header]

    # encontra coluna da sigla
    col_sigla = next(
        (c for c in df_bloco.columns if "sub" in norm(c) and len(norm(c)) <= 4),
        None
    )
    if col_sigla is None:
        # fallback: primeira coluna com "sub"
        col_sigla = next((c for c in df_bloco.columns if "sub" in norm(c)), None)
    if col_sigla is None:
        raise ValueError("Coluna de sigla não encontrada no CONSEMAVI.")

    # encontra a coluna do mês solicitado (compara normalizado)
    mes_nome = ABREV_PARA_MES[abrev_mes]          # ex: "março"
    col_mes = next(
        (c for c in df_bloco.columns if norm(str(c)) == norm(mes_nome)),
        None
    )
    if col_mes is None:
        raise ValueError(
            f"Coluna '{mes_nome}' não encontrada no bloco {ano_ref} do CONSEMAVI.\n"
            f"Colunas disponíveis: {[str(c) for c in df_bloco.columns]}"
        )

    df_bloco = df_bloco[[col_sigla, col_mes]].copy()
    df_bloco.columns = ["sigla", "consemavi"]
    df_bloco["sigla"] = df_bloco["sigla"].astype(str).str.strip().apply(norm_sigla)
    df_bloco = df_bloco[df_bloco["sigla"].str.match(r"^[a-z]{1,3}$", na=False)]
    df_bloco["consemavi"] = df_bloco["consemavi"].apply(to_float)

    return df_bloco.groupby("sigla", as_index=False)["consemavi"].sum()


# ── Montar DataFrame final ────────────────────────────────────────────────────
def montar_df_final(
    caminho_word: Optional[Path],
    caminho_consemavi: Optional[Path],
    ano_ref: int,
    abrev_mes_manual: Optional[str] = None,
) -> tuple[pd.DataFrame, str]:
    """
    Retorna (df_final, abrev_mes).

    Funciona em três situações:
      1. Apenas Word CONVIAS
      2. Apenas Excel CONSEMAVI
      3. Word + Excel juntos

    Regra do mês:
      - Se tiver Word, o mês é descoberto pelo Word.
      - Se tiver só Excel, o mês vem do painel do Streamlit.
    """
    if caminho_word is None and caminho_consemavi is None:
        raise ValueError("Envie pelo menos um arquivo: Word CONVIAS ou Excel CONSEMAVI.")

    # Cria tabelas vazias para evitar erro quando um dos arquivos não for enviado
    df_convias = pd.DataFrame(columns=["sigla", "convias"])
    df_consemavi = pd.DataFrame(columns=["sigla", "consemavi"])

    abrev_mes = None

    # Se tiver Word, lê CONVIAS e descobre o mês pelo próprio Word
    if caminho_word is not None:
        df_convias, abrev_mes = ler_word_convias(caminho_word)

    # Se tiver Excel, lê CONSEMAVI
    if caminho_consemavi is not None:

        # Se não veio Word, então ainda não sabemos o mês.
        # Nesse caso, usamos o mês escolhido no painel do Streamlit.
        if abrev_mes is None:
            abrev_mes = abrev_mes_manual

        # Se mesmo assim não tiver mês, o sistema não sabe qual coluna do Excel pegar.
        if abrev_mes is None:
            raise ValueError(
                "Não foi possível identificar o mês. "
                "Quando enviar apenas o Excel CONSEMAVI, selecione o mês no painel."
            )

        df_consemavi = ler_excel_consemavi(
            caminho_excel=caminho_consemavi,
            ano_ref=ano_ref,
            abrev_mes=abrev_mes,
        )

    # Junta os dados.
    # Se tiver só Word, vem CONVIAS e CONSEMAVI fica zero.
    # Se tiver só Excel, vem CONSEMAVI e CONVIAS fica zero.
    # Se tiver os dois, junta os dois pela sigla.
    df = pd.merge(df_convias, df_consemavi, on="sigla", how="outer")

    df["convias"] = df["convias"].fillna(0).apply(to_float)
    df["consemavi"] = df["consemavi"].fillna(0).apply(to_float)

    df["total_requalificado"] = df["convias"] + df["consemavi"]

    return df.sort_values("sigla").reset_index(drop=True), abrev_mes


# ── Escrever no Excel modelo ───────────────────────────────────────────────────
def preencher_modelo(
    df: pd.DataFrame,
    abrev_mes: str,
    caminho_modelo: Path,
    caminho_saida: Path,
) -> Path:
    """
    Preenche o modelo Acompanhamento_mensal_PDM_-_2026.xlsx (aba Meta 49).

    Estrutura do modelo:
      Linha 2 (idx): cabeçalho  →  Subprefeitura | Sigla | Indicador | Jan … Dez | Total
      A cada 3 linhas por subprefeitura:
        - Total requalificado (m2)
        - SMSUB/CONVIAS
        - SMSUB/CONSEMAVI
      Linha 99: TOTAL:
    """
    col_idx = ABREV_PARA_COL[abrev_mes]   # coluna (0-based) onde escrever

    wb = openpyxl.load_workbook(caminho_modelo)
    ws = wb["Meta 49"]

    # Mapeia sigla → (linha_total, linha_convias, linha_consemavi)  [1-based]
    mapa: dict[str, tuple[int, int, int]] = {}
    rows_ws = list(ws.iter_rows(values_only=True))

    sigla_atual = None
    linhas_bloco: list[int] = []

    for r_idx, row in enumerate(rows_ws):
        sigla_cel = row[1]  # coluna B = Sigla
        ind_cel = norm(row[2]) if row[2] else ""

        if sigla_cel and norm(sigla_cel) not in {"sigla", ""}:
            s = norm_sigla(sigla_cel)
            if s != sigla_atual:
                sigla_atual = s
                linhas_bloco = []

        if sigla_atual:
            if "total requalificado" in ind_cel:
                linhas_bloco = [r_idx + 1]  # 1-based
            elif "convias" in ind_cel and len(linhas_bloco) == 1:
                linhas_bloco.append(r_idx + 1)
            elif "consemavi" in ind_cel and len(linhas_bloco) == 2:
                linhas_bloco.append(r_idx + 1)
                mapa[sigla_atual] = tuple(linhas_bloco)  # type: ignore

    def escrever(linha_1based: int, col_0based: int, valor):
        c = col_0based + 1  # openpyxl usa 1-based
        cell = ws.cell(linha_1based, c)
        if not isinstance(cell, MergedCell):
            cell.value = valor

    # Preenche dados por sigla
    for _, row in df.iterrows():
        sigla = norm_sigla(row["sigla"])
        if sigla not in mapa:
            continue
        l_total, l_conv, l_cons = mapa[sigla]
        escrever(l_total, col_idx, row["total_requalificado"])
        escrever(l_conv,  col_idx, row["convias"])
        escrever(l_cons,  col_idx, row["consemavi"])

    # Recalcula totais de linha (soma de todos os meses)
    for sigla, (l_total, l_conv, l_cons) in mapa.items():
        for lin in (l_total, l_conv, l_cons):
            soma = sum(
                (ws.cell(lin, c + 1).value or 0)
                for c in ABREV_PARA_COL.values()
                if not isinstance(ws.cell(lin, c + 1), MergedCell)
            )
            escrever(lin, COL_TOTAL, soma if soma else None)

    # Recalcula linha TOTAL: (linha 100 = row index 99, 1-based = 100)
    linha_total_wb = None
    for r_idx, row in enumerate(rows_ws):
        if norm(row[0]) == "total:":
            linha_total_wb = r_idx + 1
            break

    if linha_total_wb:
        for col_0 in list(ABREV_PARA_COL.values()) + [COL_TOTAL]:
            soma_col = 0
            for sigla, (l_total, _, _) in mapa.items():
                v = ws.cell(l_total, col_0 + 1).value
                soma_col += v if isinstance(v, (int, float)) else 0
            escrever(linha_total_wb, col_0, soma_col if soma_col else None)

    wb.save(caminho_saida)
    return caminho_saida


# ── Entry point ───────────────────────────────────────────────────────────────

def gerar_relatorio(
    caminho_word: Optional[Path] = None,
    caminho_consemavi: Optional[Path] = None,
    ano_ref: int = 2026,
    caminho_modelo: Optional[Path] = None,
    caminho_saida: Optional[Path] = None,
    abrev_mes_manual: Optional[str] = None,
) -> tuple[Path, str]:
    """
    Gera o Excel preenchido.

    Aceita:
      - somente Word CONVIAS;
      - somente Excel CONSEMAVI;
      - Word + Excel juntos.

    Retorna (caminho_saida, abrev_mes).
    """
    caminho_modelo = caminho_modelo or arquivo_modelo_padrao

    if caminho_word is None and caminho_consemavi is None:
        raise FileNotFoundError("Envie pelo menos um arquivo: Word CONVIAS ou Excel CONSEMAVI.")

    if caminho_word is not None and not caminho_word.exists():
        raise FileNotFoundError(f"Word CONVIAS não encontrado: {caminho_word}")

    if caminho_consemavi is not None and not caminho_consemavi.exists():
        raise FileNotFoundError(f"Excel CONSEMAVI não encontrado: {caminho_consemavi}")

    if not caminho_modelo.exists():
        raise FileNotFoundError(f"Modelo não encontrado: {caminho_modelo}")

    df, abrev_mes = montar_df_final(
        caminho_word=caminho_word,
        caminho_consemavi=caminho_consemavi,
        ano_ref=ano_ref,
        abrev_mes_manual=abrev_mes_manual,
    )

    if caminho_saida is None:
        caminho_saida = exit_dir / f"meta49_{ano_ref}_{abrev_mes}.xlsx"

    preencher_modelo(df, abrev_mes, caminho_modelo, caminho_saida)

    return caminho_saida, abrev_mes